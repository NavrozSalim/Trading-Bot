"""One Excel row per trade, updated when the position closes."""

from __future__ import annotations

import os
from datetime import datetime, timezone
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from trading_bot.monitoring.logger import get_logger

log = get_logger("trading_bot.trade_excel")

HEADERS = [
    "Entry time",
    "Exit time",
    "Symbol",
    "Timeframe",
    "Buy or sell",
    "Setup",
    "Setup low",
    "Setup high",
    "Sweep price",
    "Breakout price",
    "Entry type",
    "Entry price",
    "Stop",
    "Target",
    "Points",
    "Balance",
    "Risk percent",
    "Risk amount",
    "Lot",
    "Ticket",
    "Exit price",
    "Profit or loss",
    "Exit reason",
    "Mode",
]

_TICKET_COL = HEADERS.index("Ticket") + 1


def with_clock(value: object, clock: object) -> str:
    """Price or name plus the chart time, such as '4151.561 07:21'."""
    text = "" if value is None else str(value)
    when = "" if clock is None else str(clock)
    return f"{text} {when}".strip()


def _exit_label(exit_price: float | None, stop: object, target: object, fallback: str) -> str:
    if exit_price is None:
        return fallback
    try:
        if stop not in ("", None) and abs(float(exit_price) - float(stop)) <= 0.05:
            return "stop"
        if target not in ("", None) and abs(float(exit_price) - float(target)) <= 0.05:
            return "target"
    except (TypeError, ValueError):
        return fallback
    return fallback


class TradeExcel:
    def __init__(self, path: Path) -> None:
        self.path = path

    def record_open(self, row: dict[str, object]) -> None:
        book = self._book()
        values = [row.get(name, "") for name in HEADERS]
        book.active.append(values)
        self._save(book)

    def record_close(
        self,
        ticket: str,
        *,
        exit_price: float | None,
        profit: float | None,
        exit_reason: str,
        exit_time: datetime | None = None,
    ) -> None:
        book = self._book()
        sheet = book.active
        stamp = (exit_time or datetime.now(timezone.utc)).strftime("%Y-%m-%d %H:%M:%S")
        for cells in sheet.iter_rows(min_row=2):
            if str(cells[_TICKET_COL - 1].value or "") != str(ticket):
                continue
            cells[HEADERS.index("Exit time")].value = stamp
            cells[HEADERS.index("Exit price")].value = exit_price
            cells[HEADERS.index("Profit or loss")].value = profit
            cells[HEADERS.index("Exit reason")].value = _exit_label(
                exit_price,
                cells[HEADERS.index("Stop")].value,
                cells[HEADERS.index("Target")].value,
                exit_reason,
            )
            self._save(book)
            return
        book.close()
        log.warning("trade_excel_ticket_missing", ticket=ticket)

    def _book(self) -> Workbook:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists() and self.path.stat().st_size > 0:
            try:
                return load_workbook(self.path)
            except (BadZipFile, InvalidFileException, EOFError, KeyError, ValueError) as exc:
                log.warning("trade_excel_unreadable", path=str(self.path), error=str(exc))
                self._park_broken_file()
        return self._new_book()

    def _park_broken_file(self) -> None:
        stamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
        broken = self.path.with_name(f"{self.path.stem}-broken-{stamp}{self.path.suffix}")
        try:
            self.path.replace(broken)
            log.warning("trade_excel_parked", path=str(broken))
        except OSError as exc:
            log.warning("trade_excel_broken_not_moved", path=str(self.path), error=str(exc))

    @staticmethod
    def _new_book() -> Workbook:
        book = Workbook()
        sheet: Worksheet = book.active
        sheet.title = "Trades"
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        return book

    def _save(self, book: Workbook) -> None:
        """Write a finished copy, then replace the real file in one step.

        Saving straight onto trades.xlsx empties it first. Excel then opens a
        half-written file and reports that the format is not valid.
        """
        temporary = self.path.with_name(f".{self.path.name}.tmp")
        try:
            book.save(temporary)
            os.replace(temporary, self.path)
        except PermissionError:
            log.warning("trade_excel_locked", path=str(self.path))
        except OSError as exc:
            log.warning("trade_excel_save_failed", path=str(self.path), error=str(exc))
        finally:
            if temporary.exists():
                try:
                    temporary.unlink()
                except OSError:
                    pass
            book.close()
