from openpyxl import load_workbook

from trading_bot.trading import trade_excel
from trading_bot.trading.trade_excel import TradeExcel, with_clock


def test_a_price_keeps_its_chart_time() -> None:
    assert with_clock(4151.561, "07:21") == "4151.561 07:21"
    assert with_clock("yellow", "07:20") == "yellow 07:20"
    assert with_clock("4154.416", "") == "4154.416"


def test_trade_excel_saves_open_and_close(tmp_path) -> None:  # type: ignore[no-untyped-def]
    book = TradeExcel(tmp_path / "trades.xlsx")
    book.record_open(
        {
            "Entry time": "2026-09-25 08:00:00",
            "Symbol": "XAUUSD",
            "Timeframe": "1M",
            "Buy or sell": "BUY",
            "Setup": "yellow",
            "Setup low": "4398.2",
            "Setup high": "4400",
            "Sweep price": "4398.2",
            "Breakout price": "4405",
            "Entry type": "candle close",
            "Entry price": 4405,
            "Stop": 4397.4,
            "Target": 4411.8,
            "Points": 680,
            "Balance": 500,
            "Risk percent": 2,
            "Risk amount": 10,
            "Lot": 0.02,
            "Ticket": "1001",
            "Mode": "DEMO",
        }
    )
    book.record_close("1001", exit_price=4411.8, profit=12.5, exit_reason="closed")

    sheet = load_workbook(book.path).active
    assert sheet["E2"].value == "BUY"
    assert sheet["S2"].value == 0.02
    assert sheet["T2"].value == "1001"
    assert sheet["U2"].value == 4411.8
    assert sheet["V2"].value == 12.5
    assert sheet["W2"].value == "target"
    assert not (tmp_path / ".trades.xlsx.tmp").exists()


def test_a_broken_workbook_is_replaced_and_the_old_file_is_kept(tmp_path) -> None:  # type: ignore[no-untyped-def]
    path = tmp_path / "trades.xlsx"
    path.write_bytes(b"this is not an excel file")
    TradeExcel(path).record_open({"Symbol": "XAUUSD", "Buy or sell": "BUY", "Ticket": "9"})
    sheet = load_workbook(path).active
    assert sheet["C2"].value == "XAUUSD"
    parked = list(tmp_path.glob("trades-broken-*.xlsx"))
    assert len(parked) == 1
    assert parked[0].read_bytes() == b"this is not an excel file"


def test_a_locked_workbook_is_left_untouched(tmp_path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    path = tmp_path / "trades.xlsx"
    TradeExcel(path).record_open({"Symbol": "XAUUSD", "Ticket": "1"})
    before = path.read_bytes()

    def locked(src: str, dst: str) -> None:
        raise PermissionError("Excel has the file open")

    monkeypatch.setattr(trade_excel.os, "replace", locked)
    TradeExcel(path).record_open({"Symbol": "XAUUSD", "Ticket": "2"})
    assert path.read_bytes() == before
    assert not (tmp_path / ".trades.xlsx.tmp").exists()
